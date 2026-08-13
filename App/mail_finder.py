import re
try:
    import pymupdf as fitz
except ImportError:
    import fitz
from docx import Document


EMAIL_PATTERN = re.compile(
    r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b"
)


def extract_text_from_pdf(file):
    text = ""
    file_bytes = file.read()
    with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
        for page in pdf:
            text += page.get_text()
    return text


def extract_text_from_docx(file):
    if hasattr(file, 'path') and file.path:
        document = Document(file.path)
    else:
        document = Document(file)

    text = []

    # Normal paragraphs
    for paragraph in document.paragraphs:
        text.append(paragraph.text)

    # Tables
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text.append(cell.text)

    return "\n".join(text)


def find_email_from_resume(file):
    if not file:
        return []

    if hasattr(file, 'open'):
        try:
            file.open('rb')
        except Exception:
            pass

    if hasattr(file, 'seek'):
        file.seek(0)

    filename = getattr(file, 'name', '').lower()

    text = ""
    try:
        if filename.endswith(".pdf"):
            text = extract_text_from_pdf(file)
        elif filename.endswith(".docx"):
            text = extract_text_from_docx(file)
        else:
            return []
    except Exception as e:
        print(f"Error reading resume file {filename}: {e}")
        return []
    finally:
        if hasattr(file, 'seek'):
            file.seek(0)

    emails = EMAIL_PATTERN.findall(text)

    # Remove duplicates while preserving order
    emails = list(dict.fromkeys(emails))

    return emails